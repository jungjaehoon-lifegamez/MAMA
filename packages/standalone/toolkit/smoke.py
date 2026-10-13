"""Host install/opt-in test: real CJK text must survive office rendering with embedded fonts."""
import sys

KOREAN = "문서 도구 확인"
JAPANESE = "日本語の文書確認"

if sys.argv[1] == "create":
    from pptx import Presentation
    from pptx.util import Inches, Pt

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3))
    for text, font in [(KOREAN, "Noto Sans CJK KR"), (JAPANESE, "Noto Serif CJK JP")]:
        paragraph = box.text_frame.add_paragraph()
        run = paragraph.add_run()
        run.text = text
        run.font.name = font
        run.font.size = Pt(28)
    deck.save(sys.argv[2])
elif sys.argv[1] == "check":
    import pymupdf

    with pymupdf.open(sys.argv[2]) as pdf:
        text = "".join(page.get_text() for page in pdf)
        assert KOREAN in text and JAPANESE in text, "CJK text missing from rendered PDF"
        fonts = [font for page in pdf for font in page.get_fonts(full=True)]
        assert any("Noto" in font[3] and "CJK" in font[3] and
                   font[0] > 0 and pdf.extract_font(font[0])[3] for font in fonts), \
            "Rendered PDF has no embedded Noto CJK font"
else:
    raise ValueError("Use create or check")
